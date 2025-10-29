import os
import openpyxl
from datetime import datetime
from PySide6.QtWidgets import (
    QWidget, QVBoxLayout, QHBoxLayout, QLabel, QPushButton,
    QTableWidget, QTableWidgetItem, QTabWidget, QHeaderView,
    QMessageBox, QFileDialog, QSizePolicy, QLineEdit
)
from PySide6.QtCore import Qt, QTimer, QDir
from PySide6.QtWidgets import QStyle, QApplication

# ----------------- Cấu hình -----------------
DEBOUNCE_DELAY_MS = 500
COL_G_HEADER = "GIÁ TRỊ NGHIỆM THU THEO TIẾN ĐỘ"
COL_H_HEADER = "KH THANH TOÁN TẠM ỨNG"
COL_J_HEADER = "CÒN PHẢI THU"


# --- Excel Editor widget ---
class ExcelEditor(QWidget):
    def __init__(self, excel_path: str):
        super().__init__()
        self.setStyleSheet("QWidget { background-color: #f7f9fc; }")
        self.excel_path = excel_path
        layout = QVBoxLayout(self)
        layout.setContentsMargins(6, 6, 6, 6)

        # Nội bộ
        self.sheets_tables = {}
        self.sheets_names = []
        self.col_map = {}
        self.debounce_timer = QTimer()
        self.debounce_timer.setSingleShot(True)
        self.debounce_timer.timeout.connect(self._perform_calculation)
        self.pending_calc_row = -1

        # ---------- Thanh công cụ ----------
        control_row = QHBoxLayout()
        control_row.setSpacing(8)

        # Thêm ô tìm kiếm vào cùng hàng với các button
        self.search_input = QLineEdit()
        self.search_input.setPlaceholderText("🔍 Tìm kiếm trong sheet hiện tại...")
        self.search_input.setFixedHeight(32)
        self.search_input.setFixedWidth(400)
        self.search_input.setStyleSheet("""
            QLineEdit {
                border: 1px solid #ccc;
                border-radius: 6px;
                padding: 4px 8px;
                font-size: 9pt;
            }
            QLineEdit:focus {
                border: 1px solid #0078d4;
            }
        """)
        self.search_input.textChanged.connect(self._search_in_table)
        control_row.addWidget(self.search_input)

        control_row.addStretch()

        self.btn_change_file = QPushButton("Thay Đổi File")
        self.btn_change_file.setIcon(QApplication.style().standardIcon(QStyle.SP_FileDialogStart))
        self.btn_change_file.setFixedWidth(120)
        self.btn_change_file.clicked.connect(self._change_excel_file)
        control_row.addWidget(self.btn_change_file)

        self.btn_undo = QPushButton("Hủy Thay Đổi")
        self.btn_undo.setIcon(QApplication.style().standardIcon(QStyle.SP_DialogResetButton))
        self.btn_undo.setFixedWidth(120)
        self.btn_undo.clicked.connect(self._confirm_and_undo)
        control_row.addWidget(self.btn_undo)

        self.btn_save_excel = QPushButton("Lưu Excel")
        self.btn_save_excel.setIcon(QApplication.style().standardIcon(QStyle.SP_DialogSaveButton))
        self.btn_save_excel.setFixedWidth(120)
        self.btn_save_excel.clicked.connect(self._confirm_and_save)
        control_row.addWidget(self.btn_save_excel)

        layout.addLayout(control_row)

        # ---------- Tab hiển thị sheet ----------
        self.tab_widget = QTabWidget()
        self.tab_widget.setStyleSheet("QTabBar::tab { font-size: 9pt; }")
        layout.addWidget(self.tab_widget)

        # ---------- Label thông báo đưa xuống dưới bảng ----------
        self.msg_label = QLabel("")
        self.msg_label.setStyleSheet("color: #333; font-style: italic; padding-top: 4px;")
        layout.addWidget(self.msg_label)

        # Load file lần đầu
        self.load_excel()

    # ---------------- Search handler ----------------
    def _search_in_table(self):
        """Tìm kiếm text trong sheet hiện tại và highlight các ô/ dòng khớp."""
        query = self.search_input.text().strip().lower()
        table = self.tab_widget.currentWidget()
        if not isinstance(table, QTableWidget):
            return

        for r in range(table.rowCount()):
            for c in range(table.columnCount()):
                item = table.item(r, c)
                if item:
                    item.setBackground(Qt.white)

        if not query:
            self.msg_label.setText("")
            return

        match_count = 0
        for r in range(table.rowCount()):
            row_text = " ".join(
                (table.item(r, c).text() if table.item(r, c) else "").lower()
                for c in range(table.columnCount())
            )
            if query in row_text:
                match_count += 1
                for c in range(table.columnCount()):
                    item = table.item(r, c)
                    if item:
                        item.setBackground(Qt.yellow)

        if match_count:
            self.msg_label.setText(f"🔎 Tìm thấy {match_count} dòng khớp với '{self.search_input.text().strip()}'")
        else:
            self.msg_label.setText("Không tìm thấy kết quả.")

    # ------------------ Các hàm cũ giữ nguyên ------------------
    def _change_excel_file(self):
        new_path, _ = QFileDialog.getOpenFileName(
            self,
            "Chọn File Excel Mới",
            os.path.dirname(self.excel_path) if self.excel_path else QDir.homePath(),
            "Excel Files (*.xlsx *.xls)"
        )
        if new_path:
            if os.path.normcase(new_path) == os.path.normcase(self.excel_path):
                self.msg_label.setText("Đã chọn lại file hiện tại.")
                return
            self.excel_path = new_path
            self.load_excel()
            self.msg_label.setText(f"Đã tải file Excel mới: {os.path.basename(new_path)}")

    def _confirm_and_save(self):
        msg_box = QMessageBox(self)
        msg_box.setWindowTitle("Xác nhận")
        msg_box.setText("Bạn muốn lưu thay đổi và ghi đè lên file Excel gốc?")
        btn_yes = msg_box.addButton("Có", QMessageBox.YesRole)
        btn_no = msg_box.addButton("Không", QMessageBox.NoRole)
        msg_box.setDefaultButton(btn_no)
        msg_box.exec()
        if msg_box.clickedButton() == btn_yes:
            if self.save_overwrite():
                self.load_excel()

    def _confirm_and_undo(self):
        current_table_widget = self.tab_widget.currentWidget()
        if not isinstance(current_table_widget, QTableWidget):
            self.msg_label.setText("Không có bảng tính để hủy thay đổi.")
            return
        msg_box = QMessageBox(self)
        msg_box.setWindowTitle("Xác nhận")
        msg_box.setText("Bạn có chắc chắn muốn hủy các thay đổi?")
        btn_yes = msg_box.addButton("Có", QMessageBox.YesRole)
        btn_no = msg_box.addButton("Không", QMessageBox.NoRole)
        msg_box.setDefaultButton(btn_no)
        msg_box.exec()
        if msg_box.clickedButton() == btn_yes:
            self._undo_changes(current_table_widget)

    def _undo_changes(self, table: QTableWidget):
        sheet_name = self.tab_widget.tabText(self.tab_widget.currentIndex())
        if not hasattr(table, 'original_sheet_data'):
            self.msg_label.setText("Không có dữ liệu gốc để khôi phục.")
            return
        original_data = table.original_sheet_data
        data_rows = original_data['data_rows']
        table.setRowCount(len(data_rows))
        if 'header_labels' in original_data:
            table.setColumnCount(len(original_data['header_labels']))
        for ri, row in enumerate(data_rows):
            for ci, val in enumerate(row):
                table.setItem(ri, ci, QTableWidgetItem(val))
        self.msg_label.setText(f"Đã hủy thay đổi trong sheet '{sheet_name}'.")

    def _perform_calculation(self):
        if self.pending_calc_row == -1:
            return
        row = self.pending_calc_row
        table = self.tab_widget.currentWidget()
        sheet_name = self.tab_widget.tabText(self.tab_widget.currentIndex())
        if sheet_name not in self.col_map:
            self.pending_calc_row = -1
            return
        col_indices = self.col_map[sheet_name]
        g_idx, h_idx, j_idx = col_indices.get('G', -1), col_indices.get('H', -1), col_indices.get('J', -1)
        if g_idx == -1 or h_idx == -1 or j_idx == -1:
            self.pending_calc_row = -1
            return
        try:
            def to_float(item):
                if not item or not item.text().strip():
                    return 0.0
                return float(item.text().replace(',', '').strip())
            val_g = to_float(table.item(row, g_idx))
            val_h = to_float(table.item(row, h_idx))
            con_phai_thu = val_g - val_h
            item_j = QTableWidgetItem(f"{con_phai_thu:,.2f}")
            item_j.setFlags(Qt.ItemIsSelectable | Qt.ItemIsEnabled)
            table.setItem(row, j_idx, item_j)
            self.msg_label.setText(f"✅ Đã tính toán: CÒN PHẢI THU={con_phai_thu:,.2f} tại dòng {row+1}")
        except Exception as e:
            self.msg_label.setText(f"Lỗi tính toán tại dòng {row+1}: {e}")
        self.pending_calc_row = -1

    def _on_cell_changed(self, row, column):
        current_table = self.tab_widget.currentWidget()
        sheet_name = self.tab_widget.tabText(self.tab_widget.currentIndex())
        if sheet_name in self.col_map:
            col_indices = self.col_map[sheet_name]
            g_idx = col_indices.get('G', -1)
            h_idx = col_indices.get('H', -1)
            if g_idx != -1 and h_idx != -1:
                if column == g_idx or column == h_idx:
                    self.pending_calc_row = row
                    self.debounce_timer.start(DEBOUNCE_DELAY_MS)

    def load_excel(self):
        self.sheets_tables.clear()
        self.sheets_names.clear()
        self.col_map.clear()
        while self.tab_widget.count() > 0:
            self.tab_widget.removeTab(0)
        self.msg_label.setText("")
        if not os.path.exists(self.excel_path):
            self.msg_label.setText(f"Không tìm thấy: {self.excel_path}")
            return
        try:
            wb = openpyxl.load_workbook(self.excel_path, data_only=True)
            for name in wb.sheetnames:
                ws = wb[name]
                all_rows_raw = list(ws.iter_rows(values_only=True))
                if not all_rows_raw:
                    continue
                data_rows = []
                header_labels = []
                max_col = 0
                header_row_index = -1
                for r_idx, row in enumerate(all_rows_raw):
                    cleaned_row = ["" if v is None else str(v).strip() for v in row]
                    if any(c for c in cleaned_row if c):
                        if header_row_index == -1:
                            header_labels = cleaned_row
                            header_row_index = r_idx
                            max_col = len(header_labels)
                            if max_col == 0:
                                max_col = len(row)
                                header_labels = cleaned_row
                        else:
                            data_rows.append(cleaned_row[:max_col])
                if not header_labels:
                    continue
                table = QTableWidget()
                table.original_sheet_data = {'header_labels': header_labels, 'data_rows': data_rows}
                table.verticalHeader().setVisible(False)
                table.setStyleSheet("font-size: 9pt;")
                table.setRowCount(len(data_rows))
                table.setColumnCount(max_col)
                table.setHorizontalHeaderLabels(header_labels)
                header = table.horizontalHeader()
                header.setStyleSheet("""
                    QHeaderView::section { 
                        background-color: #f0f0f0; 
                        color: #333; 
                        padding: 6px; 
                        border: 1px solid #ddd;
                        font-weight: bold;
                        font-size: 9pt;
                    }
                """)
                g_idx = h_idx = j_idx = -1
                for col_idx, label in enumerate(header_labels):
                    label_upper = label.strip().upper()
                    if COL_G_HEADER in label_upper:
                        g_idx = col_idx
                    elif COL_H_HEADER in label_upper:
                        h_idx = col_idx
                    elif COL_J_HEADER in label_upper:
                        j_idx = col_idx
                self.col_map[name] = {'G': g_idx, 'H': h_idx, 'J': j_idx}
                for ri, row in enumerate(data_rows):
                    for ci, val in enumerate(row):
                        if ci < max_col:
                            if ci == j_idx:
                                item = QTableWidgetItem(val)
                                item.setFlags(Qt.ItemIsSelectable | Qt.ItemIsEnabled)
                                table.setItem(ri, ci, item)
                            else:
                                table.setItem(ri, ci, QTableWidgetItem(val))
                if g_idx != -1 and h_idx != -1 and j_idx != -1:
                    for r in range(table.rowCount()):
                        self._on_cell_changed(r, g_idx)
                    table.cellChanged.connect(self._on_cell_changed)
                table.horizontalHeader().setSectionResizeMode(QHeaderView.ResizeToContents)
                table.setSortingEnabled(True)
                self.tab_widget.addTab(table, name)
                self.sheets_tables[name] = table
                self.sheets_names.append(name)
        except Exception as e:
            self.msg_label.setText("Lỗi đọc Excel: " + str(e))

    def save_overwrite(self):
        if not self.sheets_tables:
            self.msg_label.setText("Không có sheet để lưu.")
            return False
        try:
            wb = openpyxl.load_workbook(self.excel_path, data_only=False)
            for name, table in self.sheets_tables.items():
                if name not in wb.sheetnames:
                    ws = wb.create_sheet(title=name)
                    header_row_in_excel = 1
                else:
                    ws = wb[name]
                    header_row_in_excel = table.original_sheet_data.get('header_index', 0) + 1
                header = table.horizontalHeader()
                header_labels = [header.model().headerData(c, Qt.Horizontal) for c in range(table.columnCount())]
                for c, label in enumerate(header_labels):
                    ws.cell(row=header_row_in_excel, column=c + 1, value=label)
                rows = table.rowCount()
                cols = table.columnCount()
                for r in range(rows):
                    for c in range(cols):
                        item = table.item(r, c)
                        excel_row = header_row_in_excel + 1 + r
                        value_to_write = item.text().replace(',', '') if item else ""
                        if isinstance(value_to_write, str) and value_to_write.replace('.', '', 1).isdigit():
                            try:
                                value_to_write = float(value_to_write)
                            except ValueError:
                                pass
                        ws.cell(row=excel_row, column=c + 1, value=value_to_write)
                if header_row_in_excel + 1 + rows < ws.max_row:
                    rows_to_delete = ws.max_row - (header_row_in_excel + 1 + rows) + 1
                    ws.delete_rows(header_row_in_excel + 1 + rows, rows_to_delete)
            os.makedirs(os.path.dirname(self.excel_path), exist_ok=True)
            wb.save(self.excel_path)
            self.msg_label.setText("💾 Đã lưu thành công.")
            return True
        except Exception as e:
            self.msg_label.setText(f"Lỗi lưu: {str(e)}")
            return False
