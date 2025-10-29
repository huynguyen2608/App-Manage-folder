# main.py
import sys
import os
import threading
import subprocess
import shutil 
from datetime import datetime
import re
from typing import List, Dict, Any

from ui.contract_template_widget import ContractTemplateWidget
from ui.progress_tracking_widget import ExcelEditor
from ui.file_manager_widget import CustomFileSystemModel, PreviewWidget, FileManagerWidget
from ui.translation_handlers import TranslationHandler, TranslationWorker, LoadingOverlay, CustomProgressDialog

from PySide6.QtWidgets import (
    QApplication, QMainWindow, QWidget, QVBoxLayout, QHBoxLayout,
    QTreeView, QFileSystemModel, QSplitter, QTextEdit, QLabel, QPushButton,
    QTabWidget, QMessageBox, QTableWidget, QTableWidgetItem, QLineEdit,
    QHeaderView, QSizePolicy, QScrollArea, QStyle, QFileDialog, QListWidget, QListWidgetItem,
    QProgressDialog, QGridLayout, QProgressBar # THAY ĐỔI: Thêm QProgressBar
)
from PySide6.QtCore import Qt, QDir, QSize, Signal, QObject, QModelIndex, QTimer, QRect, QCoreApplication, QThread, QPoint, Signal # THAY ĐỔI: Thêm QThread, QPoint
from PySide6.QtGui import QPixmap, QImage, QIcon, QColor # THAY ĐỔI: Thêm QColor

# optional libs
try:
    import fitz # PyMuPDF for PDF rendering
except Exception:
    fitz = None

# THÊM: Import googletrans
try:
    from googletrans import Translator
except Exception:
    Translator = None 


# Optional dependencies for document handling
try:
    from docx import Document
except Exception:
    Document = None

import openpyxl

# ----------------- Config -----------------
# THAY ĐỔI: Thư mục gốc được đọc từ D:\Client
ROOT_FOLDER = r"D:\Client" if os.path.exists(r"D:\Client") else (r"D:\\" if os.path.exists(r"D:\\") else QDir.rootPath())
DEFAULT_EXCEL_PATH = r"D:\Data\report.xlsx"
CONTRACT_TEMPLATE_PATH = r"D:\Template"
DEBOUNCE_DELAY_MS = 500 # Độ trễ 0.5 giây cho tính toán công nợ
# ------------------------------------------

# THAY ĐỔI: Headers cho cột công nợ (Giả định)
COL_G_HEADER = "GIÁ TRỊ NGHIỆM THU THEO TIẾN ĐỘ"
COL_H_HEADER = "KH THANH TOÁN TẠM ỨNG"
COL_J_HEADER = "CÒN PHẢI THU"


# --- Overlay Widget (Màn hình xám) ---
class LoadingOverlay(QWidget):
    def __init__(self, parent=None):
        super().__init__(parent)
        # THAY ĐỔI: Màu xám bán trong suốt
        palette = self.palette()
        palette.setColor(self.backgroundRole(), QColor(0, 0, 0, 100)) 
        self.setPalette(palette)
        self.setAutoFillBackground(True)
        self.hide()

    def resizeEvent(self, event):
        # Đảm bảo overlay luôn phủ hết parent widget
        if self.parentWidget():
            self.setGeometry(self.parentWidget().rect())
        super().resizeEvent(event)


# ---------------- Main Window ----------------
class MainWindow(QMainWindow):
    def __init__(self):
        super().__init__()
        self.setWindowTitle("Manage Folder")
        self.resize(1300, 700)

        self.translation_handler = TranslationHandler(self)
        style = self.style()
        
        # THAY ĐỔI: Đặt Icon ứng dụng từ file đã upload (image_6af425.png)
        icon_path = "folder_icon.png"
        if os.path.exists(icon_path):
            custom_folder_icon = QIcon(icon_path)
            self.setWindowIcon(custom_folder_icon)
            QApplication.setWindowIcon(custom_folder_icon)
        else:
            print(f"CẢNH BÁO: Không tìm thấy icon tùy chỉnh tại '{icon_path}'. Sử dụng icon mặc định.")
            default_folder_icon = style.standardIcon(QStyle.SP_DirIcon)
            self.setWindowIcon(default_folder_icon)
            QApplication.setWindowIcon(default_folder_icon)

        # THAY ĐỔI: CSS để đồng nhất màu nền
        self.setStyleSheet("""
            QMainWindow {
                background-color: #f7f9fc; /* Màu nền chính của cửa sổ */
            }
            QWidget, QLabel, QPushButton, QLineEdit, QTableWidget, QTreeView {
                font-size: 10pt;
            }
            QTabWidget::pane {
                border-top: 1px solid #ccc;
                background-color: #f7f9fc; /* Màu nền của nội dung tab */
            }
            QTabBar::tab { 
                font-size: 11pt;
                padding: 6px 10px;
                min-height: 30px;
                background-color: #f0f0f0; /* Màu nền tab không chọn */
            }
            QTabBar::tab:selected {
                background-color: #f7f9fc; /* Màu nền tab chọn */
                color: #004d99;
                border-bottom: 2px solid #004d99;
            }
            QPushButton {
                padding: 6px; 
            }
            #btn_refresh { 
                padding: 6px;
            }
            #lbl_filename {
                font-size: 17pt; 
                font-weight: bold;
            }
            /* THAY ĐỔI: Để list search không bị tách biệt quá nhiều */
            QListWidget { 
                border-top: none; 
            }
        """)

        tabs = QTabWidget()
        self.setCentralWidget(tabs)
        
        
        # --- Tab 1: Quản lý file ---
        t_file = QWidget()
        # THÊM: Đồng nhất màu nền cho widget chứa nội dung tab
        t_file.setStyleSheet("QWidget { background-color: #f7f9fc; }") 
        tfile_layout = QVBoxLayout(t_file)
        tfile_layout.setContentsMargins(6,6,6,6)
        
        tree_container = QWidget()
        tree_container.setStyleSheet("QWidget { background-color: #f7f9fc; }")
        self.tree_vlayout = QVBoxLayout(tree_container) 
        self.tree_vlayout.setContentsMargins(0, 0, 0, 0)
        
        # Search bar
        search_layout = QHBoxLayout()
        self.search_input = QLineEdit()
        self.search_input.setPlaceholderText("Tìm kiếm gần đúng tên file...") 
        self.search_input.setFixedHeight(30)
        
        search_icon = style.standardIcon(QStyle.SP_FileDialogContentsView)
        self.search_input.addAction(search_icon, QLineEdit.LeadingPosition)
        # THAY ĐỔI: Sử dụng textChanged thay vì returnPressed để tìm kiếm hiển thị modal
        self.search_input.textChanged.connect(self.on_search) 
        
        search_layout.addWidget(self.search_input)
        self.tree_vlayout.addLayout(search_layout) 
        
        # THAY ĐỔI: List Widget hiển thị kết quả Search (Tích hợp, không dùng Popup)
        self.search_results_list = QListWidget()
        self.search_results_list.setFocusPolicy(Qt.NoFocus) 
        self.search_results_list.setStyleSheet("QListWidget { border: 1px solid #ccc; background-color: white; }")
        # THAY ĐỔI (QUAN TRỌNG): Cố định chiều cao 160px
        self.search_results_list.setFixedHeight(160) 
        self.search_results_list.hide() # Ẩn mặc định
        
        self.search_results_list.itemClicked.connect(self._select_file_from_search)
        
        # Thêm thẳng vào layout (ngay dưới search_input)
        self.tree_vlayout.addWidget(self.search_results_list) 
        
        # NÚT THÊM FILE & NÚT XÓA FILE & NÚT THAY ĐỔI ROOT & NÚT REFRESH
        file_actions_layout = QHBoxLayout()
        
        self.btn_add_file = QPushButton("Thêm File")
        self.btn_add_file.setIcon(style.standardIcon(QStyle.SP_FileIcon)) 
        self.btn_add_file.setToolTip("Sao chép file từ máy tính vào thư mục đang chọn.")
        self.btn_add_file.clicked.connect(self.add_file_to_current_folder)
        file_actions_layout.addWidget(self.btn_add_file)
        
        self.btn_delete_file = QPushButton("Xóa File/Folder")
        self.btn_delete_file.setIcon(style.standardIcon(QStyle.SP_TrashIcon))
        self.btn_delete_file.setToolTip("Xóa mục đang chọn.")
        self.btn_delete_file.clicked.connect(self.delete_selected_file)
        file_actions_layout.addWidget(self.btn_delete_file)
        
        # THÊM: NÚT THAY ĐỔI ROOT FILE
        self.btn_change_root = QPushButton("Thay đổi Folder gốc")
        self.btn_change_root.setIcon(style.standardIcon(QStyle.SP_DirHomeIcon))
        self.btn_change_root.setToolTip("Thay đổi thư mục gốc của cây thư mục.")
        self.btn_change_root.clicked.connect(self.change_root_folder)
        file_actions_layout.addWidget(self.btn_change_root)
        
        file_actions_layout.addStretch() 
        
        self.btn_refresh = QPushButton()
        self.btn_refresh.setObjectName("btn_refresh") 
        self.btn_refresh.setIcon(style.standardIcon(QStyle.SP_BrowserReload)) 
        self.btn_refresh.setToolTip("Làm mới cây thư mục.")
        self.btn_refresh.setFixedWidth(35) 
        self.btn_refresh.clicked.connect(self.refresh_tree)
        file_actions_layout.addWidget(self.btn_refresh)
        
        self.tree_vlayout.addLayout(file_actions_layout) 
        
        # SPLITTER VÀ TREE
        splitter = QSplitter(Qt.Horizontal)
        tfile_layout.addWidget(splitter)
        
        self.model = CustomFileSystemModel()
        self.model.setRootPath(ROOT_FOLDER)
        self.model.setFilter(QDir.NoDotAndDotDot | QDir.AllEntries)
        self.tree = QTreeView()
        self.tree.setModel(self.model)
        self.tree.setRootIndex(self.model.index(ROOT_FOLDER))
        
        self.tree.setSizePolicy(QSizePolicy.Expanding, QSizePolicy.Expanding)
        self.tree.setHeaderHidden(False) 
        # THÊM: Đồng nhất màu nền cho Tree View
        self.tree.setStyleSheet("QTreeView { background-color: white; border: 1px solid #ccc; }")

        self.tree.hideColumn(1) 
        self.tree.hideColumn(2) 
        
        header = self.tree.header()
        header.setSectionResizeMode(0, QHeaderView.Stretch) 
        header.setSectionResizeMode(3, QHeaderView.ResizeToContents) 
        header.setStretchLastSection(False) 

        self.tree.clicked.connect(self.on_tree_clicked)
        self.tree_vlayout.addWidget(self.tree)
        splitter.addWidget(tree_container)

        # Preview area (right)
        right_widget = QWidget()
        right_widget.setStyleSheet("QWidget { background-color: #f7f9fc; }")
        right_layout = QVBoxLayout(right_widget)
        right_layout.setContentsMargins(0,0,0,0)
        
        header_row = QHBoxLayout()
        
        self.lbl_filename = QLabel("Chọn một tệp để xem trước")
        self.lbl_filename.setObjectName("lbl_filename")
        header_row.addWidget(self.lbl_filename)
        header_row.addStretch()
        
        self.btn_open_default = QPushButton("Mở File") 
        self.btn_open_default.setIcon(style.standardIcon(QStyle.SP_DialogOpenButton)) 
        self.btn_open_default.setEnabled(False)
        self.btn_open_default.clicked.connect(self.open_with_default_app)
        header_row.addWidget(self.btn_open_default)
        right_layout.addLayout(header_row)

        self.preview = PreviewWidget()
        right_layout.addWidget(self.preview)
        
        self.lbl_file_info = QLabel("Ngày: N/A")
        self.lbl_file_info.setStyleSheet("color: gray; font-size: 9pt; padding: 4px; border-top: 1px solid #ccc;")
        self.lbl_file_info.setAlignment(Qt.AlignRight | Qt.AlignVCenter)
        self.lbl_file_info.hide()
        right_layout.addWidget(self.lbl_file_info)

        splitter.addWidget(right_widget)
        splitter.setSizes([400, 800]) 

        tabs.addTab(t_file, style.standardIcon(QStyle.SP_DirIcon), "Quản lý file")

        # --- Tab 2: Theo dõi tiến độ (Excel editor) ---
        t_progress = QWidget()
        t_progress.setStyleSheet("QWidget { background-color: #f7f9fc; }") # Đồng nhất màu nền
        progress_layout = QVBoxLayout(t_progress)
        progress_layout.setContentsMargins(6,6,6,6)

        self.excel_editor = ExcelEditor(DEFAULT_EXCEL_PATH)
        progress_layout.addWidget(self.excel_editor)
        
        tabs.addTab(t_progress, style.standardIcon(QStyle.SP_BrowserReload), "Theo dõi tiến độ")
        
        # --- Tab 3: Mẫu hợp đồng ---
        t_contract = QWidget()
        t_contract.setStyleSheet("QWidget { background-color: #f7f9fc; }") # Đồng nhất màu nền
        contract_layout = QVBoxLayout(t_contract)
        contract_layout.setContentsMargins(0, 0, 0, 0)
        self.contract_widget = ContractTemplateWidget()
        contract_layout.addWidget(self.contract_widget)
        
        # KẾT NỐI TÍN HIỆU DỊCH TỪ WIDGET CON
        self.contract_widget.translation_start_signal.connect(self.start_translation)
        
        tabs.addTab(t_contract, style.standardIcon(QStyle.SP_FileLinkIcon), "Mẫu hợp đồng")

        self.current_file = None
        
        # THÊM: Worker and Thread attributes
        self.worker = None
        self.worker_thread = None
        
        # THÊM: Overlay và Custom Dialog
        self.overlay = LoadingOverlay(self) # Màn hình xám
        self.progress_dialog_content = CustomProgressDialog(self)
        self.progress_dialog_content.hide()
        # KẾT NỐI NÚT HỦY CỦA DIALOG VỚI HÀM HỦY
        self.progress_dialog_content.canceled.connect(self.cancel_translation)
        
        # Connection to handle window movement (Loading popup follows the app)
        self.recenter_timer = QTimer(self)
        self.recenter_timer.timeout.connect(self.recenter_loading_dialog)
        self.recenter_timer.start(100) # Kiểm tra mỗi 100ms
        
        self.resizeEvent = self.on_main_window_resize_or_move # Ghi đè resizeEvent

    # ----------------------------------------------
    # PHƯƠNG THỨC: QUẢN LÝ LOADING VÀ DI CHUYỂN
    # ----------------------------------------------
    def get_centered_pos(self, size: QSize) -> QPoint:
        # Lấy vị trí trung tâm của màn hình ứng dụng trên toàn màn hình desktop
        main_rect = self.geometry()
        main_center_x = main_rect.x() + main_rect.width() // 2
        main_center_y = main_rect.y() + main_rect.height() // 2
        
        # Tính toán vị trí góc trên bên trái của dialog
        new_x = main_center_x - size.width() // 2
        new_y = main_center_y - size.height() // 2
        
        return QPoint(new_x, new_y)

    def recenter_loading_dialog(self):
        if self.progress_dialog_content.isVisible():
            # Tính toán và di chuyển dialog
            new_pos = self.get_centered_pos(self.progress_dialog_content.size())
            self.progress_dialog_content.move(new_pos)

    def on_main_window_resize_or_move(self, event):
        # Resize overlay (Phủ hết cửa sổ chính)
        self.overlay.setGeometry(self.rect())
        # Recenter dialog
        self.recenter_loading_dialog()
        # Call original handler
        super().resizeEvent(event)
        


    # ----------------------------------------------
    # PHƯƠNG THỨC: XỬ LÝ DỊCH (TRANSLATION HANDLERS)
    # ----------------------------------------------
    def start_translation(self, file_path, save_path):
        self.translation_handler.start_translation(file_path, save_path, self.contract_widget.translator)

    def cancel_translation(self):
        self.translation_handler.cancel_translation()

    def on_translation_error(self, msg):
        pass  # Handled by TranslationHandler

    def on_translation_finished(self, success):
        if success:
            self.contract_widget.load_templates()

    # ----------------------------------------------
    # CÁC PHƯƠNG THỨC KHÁC
    # ----------------------------------------------
    def change_root_folder(self):
        # ... (giữ nguyên) ...
        global ROOT_FOLDER
        
        # Mở hộp thoại chọn thư mục
        new_root = QFileDialog.getExistingDirectory(
            self, 
            "Chọn Thư mục Gốc mới", 
            ROOT_FOLDER,
            QFileDialog.ShowDirsOnly | QFileDialog.DontResolveSymlinks
        )

        if new_root and os.path.isdir(new_root):
            if new_root == ROOT_FOLDER:
                self.statusBar().showMessage("Thư mục gốc không thay đổi.", 3000)
                return
            
            # Cập nhật biến global (chỉ dùng để tham chiếu mặc định lần sau)
            ROOT_FOLDER = new_root
            
            # Cập nhật Model và Tree View
            self.model.setRootPath(new_root)
            self.tree.setRootIndex(self.model.index(new_root))
            self.tree.collapseAll() # Thu gọn tất cả cho thư mục mới

            self.preview.show_text(f"Thư mục gốc đã được thay đổi thành:\n{new_root}")
            self.lbl_filename.setText(os.path.basename(new_root) or new_root)
            self.statusBar().showMessage(f"Thư mục gốc đã thay đổi: {new_root}", 5000)
        else:
            self.statusBar().showMessage("Không thay đổi thư mục gốc.", 3000)


    # ----------------------------------------------
    # PHƯƠNG THỨC: REFRESH TREEVIEW
    # ----------------------------------------------
    def refresh_tree(self):
        # ... (giữ nguyên) ...
        current_index = self.tree.currentIndex()
        old_root_path = self.model.rootPath()
        
        self.model.setRootPath(old_root_path) 
        self.tree.setRootIndex(self.model.index(old_root_path))
        
        if current_index.isValid():
            current_path = self.model.filePath(current_index)
            new_index = self.model.index(current_path)
            if new_index.isValid():
                self.tree.setCurrentIndex(new_index)
                self.tree.scrollTo(new_index)
                
        self.statusBar().showMessage("Đã làm mới cây thư mục", 3000)

    # ----------------------------------------------
    # PHƯƠNG THỨC: TÌM KIẾM FILE VÀ HIGHLIGHT
    # ----------------------------------------------
    def search_files_and_prepare_highlight(self, search_term: str) -> List[Dict[str, str]]:
        # ... (giữ nguyên) ...
        """Tìm kiếm file gần đúng, trả về tên file gốc và tên đã được highlight (chứa tag span)."""
        search_results = []
        search_term = search_term.strip()
        
        if not search_term:
            return []

        # Pattern: '(?i)' để tìm kiếm không phân biệt chữ hoa/thường (case-insensitive)
        # Sử dụng ROOT_FOLDER hiện tại
        current_root = self.model.rootPath()
        regex_pattern = f"(?i)({re.escape(search_term)})"
        
        for root, _, files in os.walk(current_root):
            for file_name in files:
                
                if search_term.lower() in file_name.lower():
                    
                    # Tạo chuỗi HTML để highlight.
                    highlighted_name = re.sub(
                        regex_pattern, 
                        r'<span style="background-color: yellow;">\1</span>',
                        file_name
                    )
                    
                    search_results.append({
                        "original_name": file_name, 
                        "highlighted_name": highlighted_name, 
                        "full_path": os.path.join(root, file_name)
                    })
                    
        return search_results

    def on_search(self):
        # ... (giữ nguyên) ...
        """THAY ĐỔI: Xử lý tìm kiếm và hiển thị trong QListWidget tích hợp."""
        q = self.search_input.text().strip()
        
        # Nếu ô tìm kiếm trống, ẩn danh sách
        if not q:
            self.search_results_list.hide() 
            self.statusBar().clearMessage()
            return

        results = self.search_files_and_prepare_highlight(q)
        
        # Tắt tín hiệu tạm thời để tránh kích hoạt các sự kiện không cần thiết khi dọn dẹp list
        self.search_results_list.blockSignals(True) 
        self.search_results_list.clear()
        self.search_results_list.blockSignals(False)
        
        if results:
            for item_data in results:
                item = QListWidgetItem()
                
                # Hiển thị tên file gốc (original_name)
                item.setText(item_data["original_name"]) 
                item.setData(Qt.UserRole, item_data["full_path"])
                item.setData(Qt.UserRole + 1, item_data["original_name"]) 
                self.search_results_list.addItem(item)
                
            # THAY ĐỔI: Chỉ cần hiển thị QListWidget
            self.search_results_list.show()
            
            # Giữ focus trên QLineEdit (nên giữ lại dù không còn dùng Popup)
            self.search_input.setFocus() 
            
            self.statusBar().showMessage(f"Tìm thấy {len(results)} kết quả cho '{q}'", 5000)
        else:
            # THAY ĐỔI: Ẩn QListWidget trực tiếp
            self.search_results_list.hide()
            self.search_input.setFocus() 
            self.statusBar().showMessage(f"Không tìm thấy kết quả nào cho '{q}'", 5000)

    def _select_file_from_search(self, item: QListWidgetItem):
        # ... (giữ nguyên) ...
        """Xử lý khi người dùng click vào một file trong danh sách tìm kiếm."""
        file_path = item.data(Qt.UserRole)
        if file_path and os.path.exists(file_path):
            idx = self.model.index(file_path)
            if idx.isValid():
                self.tree.setCurrentIndex(idx)
                self.tree.scrollTo(idx)
                self.on_tree_clicked(idx) # Kích hoạt hàm preview
            
            # THAY ĐỔI: Ẩn QListWidget trực tiếp
            self.search_results_list.hide()
            self.search_input.setText("") # Xóa nội dung tìm kiếm sau khi chọn


    # ----------------------------------------------
    # PHƯƠNG THỨC: XÓA FILE (Đã chỉnh sửa nút xác nhận)
    # ----------------------------------------------
    def delete_selected_file(self):
        # ... (giữ nguyên) ...
        current_index = self.tree.currentIndex()
        if not current_index.isValid():
            QMessageBox.warning(self, "Lỗi", "Vui lòng chọn một file hoặc thư mục để xóa.")
            return

        file_path = self.model.filePath(current_index)
        file_name = os.path.basename(file_path)
        is_dir = os.path.isdir(file_path)

        if not os.path.exists(file_path):
             QMessageBox.warning(self, "Lỗi", "Đường dẫn không tồn tại.")
             return
        
        parent_index = current_index.parent()
        
        type_str = "thư mục" if is_dir else "file"
        
        # THAY ĐỔI (QUAN TRỌNG): Sử dụng custom QMessageBox với nút "Có" và "Không"
        msg_box = QMessageBox(self)
        msg_box.setWindowTitle("Xác nhận Xóa")
        msg_box.setText(f"Bạn có chắc muốn xóa vĩnh viễn {type_str}:\n{file_name}?")
        
        # Custom Buttons
        btn_yes = msg_box.addButton("Có", QMessageBox.YesRole)
        btn_no = msg_box.addButton("Không", QMessageBox.NoRole)
        msg_box.setDefaultButton(btn_no) # Mặc định chọn Không
        
        msg_box.exec()
        
        if msg_box.clickedButton() != btn_yes:
            return # Người dùng chọn Không hoặc đóng hộp thoại

        # Logic xóa file
        try:
            if is_dir:
                shutil.rmtree(file_path) 
            else:
                os.remove(file_path) 
            
            self.preview.show_text(f"Đã xóa {type_str}: {file_name}")
            self.lbl_filename.setText("Chọn một tệp để xem trước")
            
            self.refresh_tree()
            self.tree.setCurrentIndex(parent_index)

            QMessageBox.information(self, "Thành công", f"Đã xóa {type_str} '{file_name}' thành công.")
            
        except Exception as e:
            QMessageBox.critical(self, "Lỗi Xóa", f"Không thể xóa {type_str}:\n{e}")
    # ----------------------------------------------
    # CÁC PHƯƠNG THỨC KHÁC
    # ----------------------------------------------
        
    def add_file_to_current_folder(self):
        # ... (giữ nguyên) ...
        current_index = self.tree.currentIndex()
        if not current_index.isValid():
            target_path = self.model.rootPath()
        else:
            selected_path = self.model.filePath(current_index)
            if os.path.isdir(selected_path):
                target_path = selected_path
            else:
                target_path = os.path.dirname(selected_path)

        if not os.path.isdir(target_path):
             QMessageBox.warning(self, "Lỗi Thư mục", f"Thư mục đích không hợp lệ: {target_path}")
             return

        file_path, _ = QFileDialog.getOpenFileName(self, 
                                                   "Chọn File để Thêm vào", 
                                                   QDir.homePath(), 
                                                   "Tất cả File (*.*)")

        if file_path:
            file_name = os.path.basename(file_path)
            destination = os.path.join(target_path, file_name)
            
            if os.path.exists(destination):
                # SỬ DỤNG QMessageBox TÙY CHỈNH (Giống xác nhận xóa)
                msg_box = QMessageBox(self)
                msg_box.setWindowTitle("Xác nhận Ghi đè")
                msg_box.setText(f"File '{file_name}' đã tồn tại trong thư mục mẫu.\nBạn có muốn ghi đè?")
                
                # Thêm nút Tùy chỉnh
                btn_yes = msg_box.addButton("Có", QMessageBox.YesRole)
                btn_no = msg_box.addButton("Không", QMessageBox.NoRole)
                msg_box.setDefaultButton(btn_no) 
                
                msg_box.exec()
                
                if msg_box.clickedButton() == btn_no:
                    return

            try:
                shutil.copy2(file_path, destination)
                self.refresh_tree()
                target_index = self.model.index(target_path)
                self.tree.expand(target_index)
                QMessageBox.information(self, "Thành công", f"Đã sao chép '{file_name}' thành công vào:\n{target_path}")
                
            except Exception as e:
                QMessageBox.critical(self, "Lỗi Sao chép", f"Không thể sao chép file:\n{e}")


    def on_tree_clicked(self, index):
        # ... (giữ nguyên) ...
        # Giữ nguyên logic cũ
        file_path = self.model.filePath(index)
        file_name = self.model.fileName(index)

        if file_name.startswith("~$"):
            self.preview.show_text("Không thể xem trước file tạm thời.")
            self.lbl_filename.setText(file_name)
            self.btn_open_default.setEnabled(False)
            self.current_file = None
            return

        date_info = "Ngày: N/A"
        self.current_file = None
        self.btn_open_default.setEnabled(False)
        self.lbl_file_info.hide()

        if os.path.exists(file_path) and not os.path.isdir(file_path):
            try:
                m_time = os.path.getmtime(file_path)
                date_info = f"Ngày lưu cuối: {datetime.fromtimestamp(m_time).strftime('%d/%m/%Y %H:%M:%S')}"
                self.lbl_file_info.show()
            except Exception:
                pass

        self.lbl_file_info.setText(date_info)

        if os.path.isdir(file_path):
            if self.tree.isExpanded(index):
                self.tree.collapse(index)
            else:
                self.tree.expand(index)
            self.lbl_filename.setText(os.path.basename(file_path) or file_path)
            self.preview.show_text("Thư mục được chọn")
            return

        self.current_file = file_path
        self.lbl_filename.setText(os.path.basename(file_path)) 
        self.btn_open_default.setEnabled(True)

        ext = os.path.splitext(file_path)[1].lower()
        try:
            if ext in (".txt", ".md", ".py", ".csv", ".log"):
                with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                    txt = f.read()
                self.preview.show_text(txt)
            elif ext in (".png", ".jpg", ".jpeg", ".bmp", ".gif"):
                self.preview.show_image(file_path)
            elif ext == ".pdf":
                with open(file_path, "rb") as f:
                    b = f.read()
                self.preview.show_pdf_bytes(b)
            elif ext == ".docx":
                if Document is None:
                    self.preview.show_text("python-docx chưa cài. Cài 'python-docx' để xem .docx.")
                else:
                    doc = Document(file_path)
                    text = "\n\n".join(p.text for p in doc.paragraphs)
                    self.preview.show_text(text)
            elif ext in (".xlsx", ".xls"):
                try:
                    wb = openpyxl.load_workbook(file_path, data_only=True)
                    ws = wb.active
                    rows = []
                    for r in ws.iter_rows(values_only=True):
                        rows.append([("" if c is None else str(c)) for c in r])
                    out = "\n".join("\t".join(row) for row in rows[:200])
                    self.preview.show_text(out)
                except Exception as e:
                    self.preview.show_text("Không thể mở Excel: " + str(e))
            else:
                self.preview.show_text("Không hỗ trợ xem trước định dạng này.")
        except Exception as e:
            self.preview.show_text("Lỗi xem trước: " + str(e))

    def open_with_default_app(self):
        # ... (giữ nguyên) ...
        if not self.current_file:
            return
        try:
            if sys.platform.startswith("win"):
                os.startfile(self.current_file)
            elif sys.platform == "darwin":
                subprocess.call(["open", self.current_file])
            else:
                subprocess.call(["xdg-open", self.current_file])
        except Exception as e:
            QMessageBox.warning(self, "Mở file", f"Không thể mở file: {e}")

def main():
    # Kiểm tra và yêu cầu cài đặt thư viện nếu cần
    if fitz is None:
        print("CẢNH BÁO: Thư viện 'PyMuPDF' (fitz) chưa được cài đặt. Không thể xem trước file PDF.")
    if Document is None:
        print("CẢNH BẢO: Thư viện 'python-docx' chưa được cài đặt. Không thể xem trước file DOCX/dịch hợp đồng.")
    if Translator is None:
        print("CẢNH BẢNG: Thư viện 'googletrans' chưa được cài đặt. Không thể dịch hợp đồng. Cài đặt: pip install googletrans==4.0.0-rc1")
    
    app = QApplication(sys.argv)
    w = MainWindow()
    w.show()
    sys.exit(app.exec())

if __name__ == "__main__":
    main()